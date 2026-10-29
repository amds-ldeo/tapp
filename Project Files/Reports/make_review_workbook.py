#!/usr/bin/env python3
"""
Build the EPMA reviewer workbook from the current EPMA TAPP (pinned by VERSION).

What this is: the review instrument for domain experts. It sits between the TAPP
(addressable per field, opaque about structure) and the narrative (readable, not
commentable). Rows stay addressable; the narrative's grouping comes back as block
headers; the tier logic moves to a preamble sheet; four comment columns say what the
reviewer is being asked.

Design decisions, 2026-09-17:
  * `Keyed By` is rendered with DEFINER FIELD NAMES, not the cross-TAPP key vocabulary.
    Verified library-wide: in all 16 TAPPs every consumer key has exactly one definer,
    so the mapping is generated, never hand-written.
  * Three structure values, not two: "One value per X", "One or more values per X"
    (multi-valued but nothing repeats over it), "List of entries per X" (a definer).
  * The structure column is SPLIT BY LEVEL. A key rooted in `sample` / `sampling unit`
    has no meaning in the procedure record, which never names samples: those domains are
    stripped for the procedure column and the cell is flagged for review.
  * `>` never appears. Containment is declared once on the definer row ("List of entries
    per Sample Name") and consumers say only "per Sampling Unit Name". Only `x` survives.
  * Rows are grouped into blocks sharing (session structure, mode set) within a section,
    so display order differs from the CSV. Item numbers are assigned in CSV order and are
    the stable reference.

Layout decisions, 2026-10-26 (supersede the phrasing and row-grouping points above):
  * Structure cells are plain-language instructions — "Give one entry for each "Monitored
    Element" (4.13)" — with list names quoted, and numbered, as references to the fields
    that define them. "Entry", not "value" (2026-10-26, EPMA team review): most items are
    not quantities, and the Instructions define an entry as one complete answer in whatever
    form the Description asks for.
  * Rows follow the TAPP's own order (GROUP_ROWS = False): grouping had begun to put
    consumers above their definers.
  * Description before Example; "Example (NOT recommendations)"; "Other comments"; session
    columns in yellow; ticks in the mode columns on light green; section headings
    "Group N: name".
  * Group 1 registration and identity fields have their review cells pre-filled "Set by
    registration requirements; no review needed", greyed and kept out of the drop-downs.
  * Tabs: Instructions, Reviewer information (six free-text questions), EPMA TAPP, and the
    traceability key, which is hidden because it is for the TAPP team, not reviewers. Instructions
    walks through what to do, what each column means, how to read the structure cells, and one
    review criterion per comment column. Needed? gained "Rename (say how in Other comments)".
  * RD's hand edit of 2026-10-26, made standard: Instructions in 14 pt with banded sections and
    the column labels in their header colours (left unprotected); comments are free text, with
    no drop-downs; Group 1's 14 registration fields moved to a "Registration Metadata" tab, with
    one placeholder row left on EPMA TAPP; a contributors question on Reviewer information.
    Restored on top of it: the two-records paragraph, the list legend, the procedure-projection
    pattern, the header-row freeze (C2) and the filter. Row heights are estimated (est_height).
  * Sheets are protected (no password): on EPMA TAPP only the comment cells are editable,
    except the pre-filled registration cells; on Reviewer information only the answers.

Regenerate after any version bump; this is a snapshot of one version, like the mockups.
Output goes to `TAPPs for review/` at the library root (2026-09-30), which holds only the
latest version: writing a new version deletes the EPMA workbooks of any other version there.
"""

import csv, os, re, collections
from review_examples import example, EXAMPLES, MODES
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

# Paths resolve from this file's own location, as build_form.py does, so the script runs
# from any working directory. The TAPP is read from the `Current TAPPs/` mirror (Rule 12);
# VERSION pins the snapshot — bump it and the output name follows.
VERSION = 'v91'
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
SRC  = os.path.join(ROOT, 'Current TAPPs', 'EPMA_TAPP_%s.csv' % VERSION)
if not os.path.exists(SRC):                      # mirror stale or mid-bump
    SRC = os.path.join(ROOT, 'EPMA', 'EPMA_TAPP_%s.csv' % VERSION)
OUT_DIR = os.path.join(ROOT, 'TAPPs for review')
OUT  = os.path.join(OUT_DIR, 'EPMA_TAPP_%s_Review_Workbook.xlsx' % VERSION)


def out_path(mode):
    if mode is None:
        return OUT
    return os.path.join(OUT_DIR, 'EPMA_TAPP_%s_Review_Workbook_%s.xlsx' % (VERSION, mode.replace(' ', '_')))


def prune_other_versions():
    """Delete EPMA workbooks of any version but VERSION; git history keeps the old ones."""
    pat = re.compile(r'^EPMA_TAPP_(v\d+)_Review_Workbook(_[A-Za-z_]+)?\.xlsx$')
    gone = []
    for fn in sorted(os.listdir(OUT_DIR)):
        m = pat.match(fn)
        if m and m.group(1) != VERSION:
            os.remove(os.path.join(OUT_DIR, fn)); gone.append(fn)
    return gone

# Row order within each section. True groups rows that repeat over the same thing in the same
# modes (the 2026-09-17 layout); False keeps the TAPP's own order (chosen for the plain-keys draft
# on 2026-10-26, because grouping had started to put consumers above their definers).
GROUP_ROWS = False

# Group 1 registration and identity fields: listed on their own "Registration Metadata" tab, with
# the judgement cells greyed and locked and only "Other comments" open (RD's edit of 2026-10-26,
# made standard). Identified by name. The four coupling fields stay on the EPMA TAPP tab.
REGISTRATION = {
    'Procedure Name', 'Technique', 'Procedure Author', 'Laboratory', 'Laboratory ID',
    'Procedure Start Date', 'Funding Source for Procedure Development', 'Procedure Reference(s)',
    'Procedure DOI', 'Session Identifier', 'Analyst', 'Analysis Start Date', 'Analysis End Date',
    'Funding Source for Analysis',
}
# Section headings read "Group N: name" (RD's own edit of 2026-10-22, made standard).
SECTION_SHORT = {'Quality Control & Uncertainty': 'QC & Uncertainty'}

# Reviewer information tab (2026-10-26). All answers are free text.
REVIEWER_QUESTIONS = [
    "Reviewer's name",
    "Reviewer's email",
    "Reviewer's institute",
    'Frequently used instrument model',
    'Frequently used data acquisition software',
    'Frequently used data processing software',
    'Other contributors to this review (names and institutes)',
]
# TAPP fields left out of the review workbooks (they stay in the TAPP). RD, 2026-10-26.
NOT_REVIEWED = {'Additional Notes'}
# Round deadline, shown in Step 4 of the Instructions tab.
DEADLINE = 'December 1, 2026'

# Domains that do not exist until a session runs.
SESSION_ONLY = {'sample', 'sampling unit'}   # definers are C=N/A (7.3.3); `combined result` retired 2026-10-26

# Singular display names for the domains, pointing at the definer field.
DISPLAY = {
    'sample':             'Sample Name',
    'sampling unit':      'Sampling Unit Name',
    'target species':     'Target Species',
    'monitored property': 'Monitored Element',
    'target material':    'Target Material',
    'reported property':  'Reported Variable',
    'standard':           'Secondary Reference Material',
}

# Containment: a domain and the domain it nests inside, read off `defines: A > B` rows.
PARENT = {'sampling unit': 'sample'}

# Fields keyed `(none)` whose own description asks for more than one value. Identified by
# reading Columns B and F; `(none)` says "nothing repeats over this", not "one value".
# These are a question for reviewers, not a settled answer.
MULTI = {
    'Instrument Model',
    'Data Processing Software(s)',
    'EDS Detector Configuration',
    'Analytical Mode',
    'EDS Acquisition Mode',
    'Coupled Dataset or Publication Reference',
    'Coupled Technique(s)',
    'Target Material',
}

# Procedure-level multiplicity. The literature assessment (15 EPMA procedures in v77)
# states these as a set of values, each qualified by the phase or material it applies to:
# Beam Current 6/15 procedures, Beam Mode 7/15, Beam Diameter 6/15, Beam Damage 4/15.
# A single procedure-level value would misrepresent them. The domain they vary over
# (phase / material type) has no key at procedure level — a question for reviewers.
# Emptied 2026-09-28 (gap 1): the beam fields are keyed `target material` from v79, so the
# per-phase values have a list of their own and no longer project to "one or more".
PROC_MULTI = set()
# Emptied 2026-09-30: both open cells settled on their projected structure, since reviewers
# judge every structure cell through "Structure OK?" anyway.
#   Pre-Analysis Imaging and Screening -> "One value per procedure". 14 of 15 v86 procedures
#     state one screening approach, never varied by phase or material (McCoy+2025_UA is N).
#   Counting Statistics Error -> "One value per Reported Variable". 0 of 15 state it, so the
#     call rests on the parallel with Detection Limit (C=Advanced, keyed reported property):
#     a procedure can give a typical figure per variable, never a per-point one.
REVIEW_NOTE = {}
NOTE_TAIL = (' The procedure record cannot name samples or analysis points, so the per-unit '
             'detail appears only in the session column. Is this what a registering '
             'laboratory would expect to enter?')

PROSE = {
 'Procedure Name': 'A short descriptive name with a version number, identifying the instrument, the technique and the scope of what is measured.',
 'Technique': 'Whether the procedure uses WDS, EDS or both.',
 'Procedure Author': 'The person or laboratory responsible for the procedure; ORCID recommended for a person.',
 'Laboratory': 'The laboratory or institution hosting the instrument.',
 'Laboratory ID': "The laboratory's persistent identifier.",
 'Procedure Start Date': 'The date this procedure configuration first went into production use.',
 'Funding Source for Procedure Development': 'The funding that supported the instrument, major upgrades and development of the procedure, with agencies and grant numbers.',
 'Procedure Reference(s)': 'Publications or technical reports describing, validating or benchmarking the procedure.',
 'Procedure DOI': "The DOI of the registered procedure followed, or 'pending'.",
 'Session Identifier': "The laboratory's own run, sequence or batch identifier for the session, as generated by the instrument or acquisition software.",
 'Analyst': 'The analyst or analysts who performed the session; ORCID recommended.',
 'Analysis Start Date': 'The date the session began.',
 'Analysis End Date': 'The date the session ended.',
 'Funding Source for Analysis': 'The funding that supported these analyses, with agencies and grant numbers.',
 'Coupled Technique(s)': "Other techniques applied to the same samples whose results are interpreted together with these, with the instrument, laboratory or purpose of each; 'None' if not coupled.",
 'Coupling Description': 'What data or context passes between the techniques, and which technique is performed first and why. Required whenever a coupled technique is named.',
 'Coupled Procedure DOI': "The registered procedure DOI of the coupled technique, a publication DOI, 'pending', or 'None'.",
 'Coupled Dataset or Publication Reference': "The identifier of the dataset or publication reporting both datasets together, 'same submission', or 'pending'.",
 'Target Material': 'The material types the procedure is designed to analyse.',
 'Sample Preparation Method': "The form in which samples are presented to the instrument and the preparation that brought them to it; 'None' where material is analysed as received.",
 'Sample Name': 'The name of each sample analysed in the session, as the laboratory records it.',
 'Sampling Unit Type': 'The kind of subdivision of a sample that one row of reported results corresponds to, giving both levels where such units nest.',
 'Sampling Unit Name': 'The label of each unit analysed and the sample it belongs to; where units are too numerous to name individually, the acquisition area they belong to.',
 'Sample Persistent Identifier': 'The persistent identifier of the sample, at the level actually analysed; IGSN recommended.',
 'Sampling Unit Selection Criteria': 'The criteria by which the units analysed are selected within a sample.',
 'Pre-Analysis Imaging and Screening': 'Imaging or screening performed before analysis to locate and select the units analysed, with the technique, instrument and settings, and how individual analyses are linked back to the images.',
 'Instrument Manufacturer': 'The manufacturer of the instrument.',
 'Instrument Model': 'The model designation, including any generation or configuration suffix.',
 'Electron Source': 'The type of electron source.',
 'Acquisition Software': 'The software that controls the instrument and acquires the data, with version number.',
 'Data Processing Software(s)': 'Every software package applied to the data after acquisition, with version numbers.',
 'WDS Spectrometer Configuration': 'The number, type and crystal range of the WDS spectrometers, with manufacturer and model.',
 'EDS Detector Configuration': 'The detector type, manufacturer, number of detector elements, active area and solid angle, window type and geometry, with multiple detectors listed separately.',
 'Analytical Mode': 'The analytical modes the procedure covers.',
 'Beam Mode': 'The beam mode for point analysis. Must be consistent with beam diameter and raster dimensions.',
 'Target Material of Sampling Unit': 'The Target Material entry each analysis point belongs to, which links it to the beam conditions registered for that material.',
 'Mapping Beam Mode': 'Whether the beam was focused or defocused during X-ray mapping.',
 'Mapping Beam Current': 'The probe current in nA used during X-ray mapping.',
 'Mapping Beam Diameter': 'The beam diameter in µm during X-ray mapping.',
 'Accelerating Voltage': 'The accelerating voltage in kV, with justification for any departure from the standard operating voltage.',
 'Beam Current': 'The probe current in nA for point analysis.',
 'Beam Diameter': 'The beam diameter in µm for point analysis.',
 'Beam Raster Dimensions': 'The width and height in µm of the area over which the beam is rastered.',
 'Beam Damage Minimization': 'The measures taken to minimise beam damage, with the beam conditions used and the phases to which they are applied.',
 'Drift Correction': 'How drift in beam current and spectrometer position is monitored and corrected during a session.',
 'Target Species': 'The elements the procedure determines. Isotopes are never separate entries.',
 'Monitored Elements': 'The elements whose X-rays are measured, each listed under the target species it is used for. Usually the same elements as Target Species, but not always: an element determined without being measured (O by stoichiometry, C as CO₂) has no entry here, and an element measured only to correct or check another (Zn, to correct its overlap on Na Kα) is listed here but not under any target species.',
 'Reported Variables and Units': 'The quantities the procedure finally reports, with their units, including intermediate quantities reported alongside final ones.',
 'X-ray Detection Method per Monitored Element': 'Whether the monitored element is measured by WDS or by EDS.',
 'X-ray Line': 'The X-ray emission line measured.',
 'Diffracting Crystal': 'The diffracting crystal.',
 'WDS Spectrometer Channel': 'The spectrometer position or positions on which the element is measured, one entry per assignment, including an element measured on more than one spectrometer with the intensities aggregated.',
 'Sequence': 'The order in which the element is acquired and, where the element suite exceeds the number of spectrometers, the acquisition pass to which it belongs.',
 'Proportional Counter / Detector': 'The detector type.',
 'WDS PHA Setting': 'The pulse height analyser mode.',
 'Peak Counting Time': 'The on-peak counting time in seconds.',
 'Background Counting Time': 'The total off-peak background counting time in seconds, summed across all background positions.',
 'Background Position(s)': 'The location of each background position relative to the peak, in millimetres or sin θ and on the high- or low-energy side, or that no off-peak positions are used.',
 'EDS Live Time per Point or Pixel': 'The live time per analysis point or per pixel in seconds.',
 'EDS Acquisition Mode': 'How the beam is positioned during EDS acquisition, listing all approaches used.',
 'Dwell Time per Pixel': 'The time the beam dwells on each pixel: one value per spectrometer assignment for WDS, and a single real (clock) time per pixel for EDS. The EDS live time is recorded separately.',
 'Step Size / Pixel Size': 'The step size in micrometres, with X and Y given separately where they differ.',
 'Map Dimensions': 'The map dimensions in pixels.',
 'Map Area': 'The physical extent of the mapped area.',
 'Stage Scan vs. Beam Scan': 'Whether maps are acquired by stage scan or by beam scan.',
 'Matrix Correction Method': 'The matrix correction algorithm, which also applies when count maps are converted to concentration maps.',
 'Mass Absorption Coefficients (MACs)': 'The database of mass absorption coefficients the correction uses.',
 'X-ray Background Correction Method': 'How the background beneath the peak is estimated and subtracted.',
 'Time-Dependent Intensity Correction': 'The time-dependent intensity correction applied.',
 'Target Species Estimation Method': 'How the concentration of the element is obtained from the measured data.',
 'Halogen Correction on Oxygen': 'Whether oxygen is adjusted for halogen substitution where oxygen is calculated by stoichiometry.',
 'WDS Dead Time Correction': 'The dead-time correction algorithm and any instrument-specific constant. No measured WDS dead time is reported.',
 'EDS Spectral Processing Type': 'How net peak intensities are extracted from the spectra before quantification.',
 'Blank Correction': 'The method and reference material used to determine and subtract a blank.',
 'Normalization / Standards-Based Correction': "Any normalisation applied beyond the primary calibration, or 'None'.",
 'Calibration Factor and Determination Method': 'Any externally calibrated factor that converts the measured quantity into the reported quantity, how that factor was determined, and its uncertainty.',
 'Procedural Blank Level': "The blank level measured in the session, together with the blank's composition where the reported quantity is a ratio.",
 'Analysis Inclusion and Rejection Criteria': 'The rules deciding which individual results contribute to a combined value: the limits a result must meet and the grounds on which one is rejected. In the dataset: flag each excluded result, with the rule that excluded it.',
 'Combination Method': 'Whether the reported variable is an individual result or combines several; if combined, the statistical model used and what the results are combined over. In the dataset: report each combined value with its uncertainty and number of results, and, for each result, the combined value it contributes to.',
 'Constants and Reference Values Used': "The physical constants and citable reference values used in data reduction, with their sources, or 'None'.",
 'Primary Calibration Standard Name': "The primary reference material against which the instrument is calibrated, its source, and a citation for the accepted values used; where quantification uses a stored library or response model, that instead. 'None' means no calibration was performed.",
 'Secondary Reference Materials': 'The reference materials measured as unknowns alongside the samples, with the source of each and a citation for its accepted values.',
 'X-ray Line Overlap Corrections Applied': 'Whether an X-ray line overlap correction is applied and, if so, the overlapping lines and the correction method.',
 'Interfering Elements': 'The elements whose lines overlap the measured peak.',
 'Interference Correction Standard': 'The reference material used to calibrate the interference correction.',
 'Detection Limit': 'The detection limit, with units, stating whether the value is a typical estimate or a session-specific measurement. Must be consistent with the stated method. In the dataset: the detection limit of each analysis, where the procedure computes one.',
 'Detection Limit Method': 'The formula or approach used to calculate the detection limit, with a citation where one exists.',
 'Analytical Precision': 'The precision of repeated measurements, as a one-sigma relative standard deviation in percent, with the number of analyses and the measured value, and whether the repeats were made within one session or across sessions.',
 'Analytical Accuracy': 'The offset between measured and accepted values, as a percent relative bias, with the source of the accepted value.',
 'Counting Statistics Error': 'How the uncertainty predicted from counting statistics is computed, including the counts on any background or blank subtracted, and its sigma level. In the dataset: that uncertainty for each analysis and reported variable.',
 'EDS Dead Time': 'The percent dead time recorded by the detector.',
 'Goodness-of-Fit or Dispersion Statistic': 'The statistic used to describe the scatter among the results contributing to a combined value, or to test whether it exceeds what analytical uncertainty predicts (for example one standard deviation of the mean, or the MSWD). In the dataset: report its value for each combined value, and any further statistic beside it, named.',
 'Additional Notes': 'Any information not captured elsewhere, including anomalies, deviations from the registered procedure and instrument modifications.',
}

# ---------------------------------------------------------------- structure phrases

def parse_key(k):
    """Return (is_definer, parent_domain_or_None, [operand domains])."""
    k = k.strip()
    if k.startswith('defines:'):
        body = k[len('defines:'):].strip()
        if ' per ' in body:
            enumerated, parent = [x.strip() for x in body.split(' per ', 1)]
            return True, parent, [enumerated]
        if '>' in body:                      # defines: sample > sampling unit
            segs = [s.strip() for s in body.split('>')]
            return True, segs[-2], [segs[-1]]
        return True, None, [body]
    if k == '(none)':
        return False, None, []
    operands = []
    for part in k.split(' x '):
        segs = [s.strip() for s in part.split('>')]
        operands.append(segs[-1])            # `>` collapses: the innermost domain wins
    return False, None, operands


def chain(domain):
    """Domain plus its containment parents, innermost first.

    Expansion follows `>` (containment) only, never `defines: A per B`. Rule 7 draws that
    line: under `>` the inner domain cannot be enumerated without the outer one, whereas
    `per` is a grouping whose entries may have no parent — Monitored Elements explicitly
    admits elements that serve no target species.
    """
    out, seen = [DISPLAY[domain]], {domain}
    d = domain
    while d in PARENT and PARENT[d] not in seen:
        d = PARENT[d]; seen.add(d); out.append(DISPLAY[d])
    return ' per '.join(out)


# Plain-language structure phrasing (standard since 2026-10-26; was "One value per X").
NOUN = {                       # singular noun for each domain, as a reviewer would say it
    'sample':             'sample',
    'sampling unit':      'analysis point',
    'target material':    'target material',
    'target species':     'target species',
    'monitored property': 'monitored element',
    'reported property':  'reported variable',
    'standard':           'secondary reference material',
    'acquisition pass':   'acquisition pass',
    'preparation step':   'preparation step',
}
DEFINES = {                    # the list a definer row establishes
    'sample':            'List every sample analysed in the session.',
    'target material':   'List the material types the procedure is designed to analyse.',
    'target species':    'List the elements the procedure determines.',
    'reported property': 'List every quantity reported, with its unit.',
    'standard':          'List the secondary reference materials run as unknowns.',
    'acquisition pass':  'List the passes the measurement is run in, each with its own configuration.',
    'preparation step':  'List the stages of sample preparation, in order.',
}


DEF_NUM = {}                   # domain -> item number of its definer row, filled by build()


def Q(domain):
    """The list's name as the preamble legend gives it, quoted and numbered, so it reads as a reference."""
    name = '"%s"' % DISPLAY.get(domain, domain)
    return '%s (%s)' % (name, DEF_NUM[domain]) if domain in DEF_NUM else name


def phrase(name, key, level):
    """Plain-language structure for one record. level is 'procedure' or 'session'. Returns (text, flagged)."""
    k = key.strip()
    record = 'procedure' if level == 'procedure' else 'session'
    if k.startswith('defines:'):
        body = k[len('defines:'):].strip()
        if ' per ' in body:                                   # defines: A per B
            a, b = [x.strip() for x in body.split(' per ', 1)]
            return ('List each %s under the %s it is used for; list one measured only to correct '
                    'or check another on its own.' % (NOUN.get(a, a), Q(b)), False)
        if '>' in body:                                       # defines: A > B
            a, b = [x.strip() for x in body.split('>')][-2:]
            if level == 'procedure' and a in SESSION_ONLY:
                return ('List the %ss.' % NOUN.get(b, b), True)
            return ('For each %s, list the %ss analysed in it.' % (Q(a), NOUN.get(b, b)), False)
        return (DEFINES.get(body, 'List the %ss.' % NOUN.get(body, body)), False)
    if k == '(none)':
        multi = name in MULTI or (level == 'procedure' and name in PROC_MULTI)
        return ('Give %s for the %s.' % ('one or more entries' if multi else 'one entry', record), False)
    if k.startswith('pair:'):
        d = k[len('pair:'):].strip()
        return ('Give one entry for each pair of %s entries.' % Q(d), False)

    parts = [[s.strip() for s in p.split('>')] for p in k.split(' x ')]   # each part: [outer, ..., inner]
    inner = [p[-1] for p in parts]
    contained = [p[:-1] for p in parts]
    dropped = [d for d in inner if level == 'procedure' and d in SESSION_ONLY]
    kept = [(d, c) for d, c in zip(inner, contained) if d not in dropped]
    if not kept:                                              # projection to one value (7.3.3)
        return ('Give one entry; each %s entry in a session inherits it.' % Q(dropped[-1]), True)
    outer = next((c[-1] for _, c in kept if c and level == 'session'), None)
    if len(kept) == 1:
        d = kept[0][0]
        if outer:
            return ('For each %s, give one entry for each of its %s entries.' % (Q(outer), Q(d)), bool(dropped))
        return ('Give one entry for each %s.' % Q(d), bool(dropped))
    (d1, _), (d2, _) = kept[0], kept[1]
    if outer:                                                 # e.g. sample > sampling unit x reported property
        return ('For each %s, give one entry for every %s at every %s entry.'
                % (Q(outer), Q(d2), Q(d1)), bool(dropped))
    return ('For each %s, give one entry for each %s.' % (Q(d1), Q(d2)), bool(dropped))


MODE_PHRASE = {
    'YYYY': 'all four modes', 'NNYY': 'WDS only', 'YYNN': 'EDS only',
    'NYNY': 'mapping only', 'YNYN': 'point analysis only',
    'NNYN': 'WDS point analysis only', 'YNNN': 'EDS point analysis only',
}

# ---------------------------------------------------------------- styling
TIER_STYLE = {'Basic': ('C00000', True), 'Advanced': ('375623', True),
              'Read-Only': ('0070C0', True), 'Editable': ('7030A0', True),
              'N/A': ('808080', False)}
HEAD_FILL    = PatternFill('solid', fgColor='44546A')
SECTION_FILL = PatternFill('solid', fgColor='C9D5EA')
BLOCK_FILL   = PatternFill('solid', fgColor='EDF1F8')
FLAG_FILL    = PatternFill('solid', fgColor='FFF2CC')
TICK_FILL    = PatternFill('solid', fgColor='C6EFCE')
GREY_FILL    = PatternFill('solid', fgColor='D9D9D9')
ASK_FILL     = PatternFill('solid', fgColor='FBE5D6')
LEGEND_FILL  = PatternFill('solid', fgColor='F2F2F2')
THIN  = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Column order keeps each record's two columns together: structure then level.
# "Session level" is this workbook's name for the TAPP's Analysis-Level Tier. One word is
# used throughout — "analysis" collides with an individual analysis point, which is the
# one thing the record is NOT.
# ---------------------------------------------------------------- columns
# Each review column sits immediately after the content column it judges and shares its
# colour: the content header is the family's dark tone, the review header the mid tone,
# and the review cells below a pale wash of it.
FAMILY = {
    'item':      dict(dark='7030A0', mid='B4A7D6', pale='EDE7F5'),
    'procedure': dict(dark='2F5597', mid='8FAADC', pale='DCE4F4'),
    'session':   dict(dark='BF8F00', mid='FFD966', pale='FFF2CC'),   # yellow since 2026-10-26
    'mode':      dict(dark='44546A', mid='44546A', pale='FFFFFF'),
    'comment':   dict(dark='833C0B', mid='F4B183', pale='FBE5D6'),
}

# No double quotes in any option: the list becomes an XLSX formula string that is itself
# double-quoted, and an embedded " makes Excel offer to repair the file.
OPT_NEEDED = ['Keep', 'Drop', 'Merge with another row', 'Split into several rows',
              'Rename (say how in Other comments)']
OPT_STRUCT = ['Yes', 'Repeats over the wrong thing', 'Should be a list',
              'Should be a single value']
OPT_PTIER  = ['Yes', 'Should be required', 'Should be optional',
              'Should not be in the procedure record']
OPT_STIER  = ['Yes', 'Should be required', 'Should be optional',
              'Should be fixed by the procedure', 'Should be adjustable']

# key, header, width, colour family, kind, dropdown options
COLSPEC = [
    ('num',      '#',                            7,  None,        'content', None),
    ('name',     'Metadata item',                31, 'item',      'content', None),
    ('needed',   'Needed?',                      17, 'item',      'review',  OPT_NEEDED),
    ('prose',    'Description',                  56, None,        'content', None),
    ('example',  'Example (NOT recommendations)', 44, None,       'content', None),
    ('pstruct',  'Structure — procedure record', 28, 'procedure', 'content', None),
    ('pstructr', 'Structure OK?',                16, 'procedure', 'review',  OPT_STRUCT),
    ('ptier',    'Procedure-Level Tier',         14, 'procedure', 'content', None),
    ('ptierr',   'Tier OK?',                     16, 'procedure', 'review',  OPT_PTIER),
    ('sstruct',  'Structure — session record',   31, 'session',   'content', None),
    ('sstructr', 'Structure OK?',                16, 'session',   'review',  OPT_STRUCT),
    ('stier',    'Session-Level Tier',           14, 'session',   'content', None),
    ('stierr',   'Tier OK?',                     16, 'session',   'review',  OPT_STIER),
]
MODE_W, COMMENT_W = 10, 38
COMMENT_COL = ('comment', 'Other comments', COMMENT_W, 'comment', 'review', None)


UNLOCKED = Protection(locked=False)


def protect(sheet):
    """Lock every cell not explicitly unlocked; no password — this prevents accidental edits,
    it is not security. Filtering, resizing and selection still work."""
    p = sheet.protection
    p.sheet = True
    p.autoFilter = False; p.formatColumns = False; p.formatRows = False
    p.selectLockedCells = False; p.selectUnlockedCells = False


def est_height(cells, size=11, minimum=15):
    """Row height for wrapped text: openpyxl cannot autofit, and Excel does not refit on open.
    cells = [(text, column width)]; roughly 1.1 characters per width unit at 11 pt."""
    lines = 1
    for text, width in cells:
        if not text:
            continue
        per = max(1, int(width * 1.1 * 11 / size))
        n = sum(max(1, -(-len(part) // per)) for part in str(text).split('\n'))
        lines = max(lines, n)
    return max(minimum, lines * size * 1.35 + 4)


def colspec(modes):
    """Full column list: the fixed columns, one per mode, then the free-text comment."""
    spec = list(COLSPEC)
    for m in modes:
        spec.append(('mode:' + m, m, MODE_W, 'mode', 'content', None))
    spec.append(COMMENT_COL)
    return spec



def load():
    rows = list(csv.reader(open(SRC, encoding='utf-8-sig')))
    h = rows[0]
    sent = h.index('Literature Assessment')
    m0 = h.index('Purpose') + 1
    modes = h[m0:sent]
    out, sec, secname, idx = [], 0, '', 0
    for r in rows[1:]:
        if not r[2].strip() and not r[3].strip():
            if r[0].strip():
                sec += 1; secname = r[0].split('. ', 1)[-1]; idx = 0
            continue
        idx += 1
        out.append(dict(sec=sec, secname=secname, num='%d.%d' % (sec, idx), name=r[0].strip(),
                        C=r[2].strip(), D=r[3].strip(), key=r[8].strip(),
                        modes=''.join(x.strip() for x in r[m0:sent])))
    return out, modes


def build(mode=None):
    fields, modes = load()
    for f in fields:               # numbered from the full TAPP, so mode copies cite the same rows
        if f['key'].startswith('defines:'):
            body = f['key'][len('defines:'):].split(' per ')[0]
            DEF_NUM[body.split('>')[-1].strip()] = f['num']
    fields = [f for f in fields if f['name'] not in NOT_REVIEWED]
    missing = [f['name'] for f in fields if f['name'] not in PROSE]
    missing += [f['name'] for f in fields if f['name'] not in EXAMPLES]
    if missing:
        raise SystemExit('no sentence or example written for: %s' % missing)
    all_modes = list(modes)
    total_fields = len(fields)
    if mode is not None:
        k = all_modes.index(mode)
        fields = [f for f in fields if f['modes'][k] == 'Y']
        modes = []                       # one mode per workbook: no mode columns

    wb = Workbook()

    # ============================== sheet 1: preamble ==============================
    lg = wb.active; lg.title = 'Instructions'
    lg.sheet_view.showGridLines = False
    # Layout follows RD's hand edit of 2026-10-26: 14 pt body, banded sections, column labels
    # in their header colours. The tab is left unprotected, as in that edit.
    WA, WB = 34, 115
    PAGE = PatternFill('solid', fgColor='F2F2F2')
    STEP = PatternFill('solid', fgColor='D6DCE4')
    BOX = Border(left=Side(style='thin'), right=Side(style='thin'),
                 top=Side(style='thin'), bottom=Side(style='thin'))
    r = 1
    def merged(txt, size, bold=False, fill=PAGE, centre=False, border=None, gap=0):
        nonlocal r
        c = lg.cell(row=r, column=1, value=txt)
        c.font = Font(bold=bold, size=size); c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical='center' if centre else 'top',
                                horizontal='center' if centre else None)
        if border:
            c.border = border; lg.cell(row=r, column=2).border = border
        lg.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        lg.row_dimensions[r].height = est_height([(txt, WA + WB)], size)
        r += 1 + gap
    def title(txt):
        merged(txt, 22, bold=True, centre=True)
    def head(txt):
        nonlocal r
        r += 1
        merged(txt, 20, bold=True, centre=True, border=BOX)
    def para(txt):
        merged(txt, 14, border=BOX)
    def pairs(items, fill=PAGE, label_size=14):
        """items: (label, text, font colour or None, label fill colour or None)."""
        nonlocal r
        for label, text, colour, lfill in items:
            a = lg.cell(row=r, column=1, value=label)
            a.font = Font(bold=True, size=label_size,
                          color='FFFFFF' if lfill else (colour or '000000'))
            a.fill = PatternFill('solid', fgColor=lfill) if lfill else fill
            a.alignment = Alignment(wrap_text=True, vertical='top'); a.border = BOX
            b = lg.cell(row=r, column=2, value=text)
            b.font = Font(size=14); b.fill = fill
            b.alignment = Alignment(wrap_text=True, vertical='top'); b.border = BOX
            lg.row_dimensions[r].height = est_height([(label, WA), (text, WB)], 14)
            if label_size > 14:
                lg.row_dimensions[r].height = max(lg.row_dimensions[r].height, label_size * 1.6)
            r += 1

    review_fields = [f for f in fields if f['name'] not in REGISTRATION]
    reg_fields = [f for f in fields if f['name'] in REGISTRATION]
    n_info = 8 + len(modes)
    mode_tag = ('_' + mode.replace(' ', '_')) if mode else ''
    return_name = 'EPMA_TAPP_%s_Review%s_<your initials>.xlsx (e.g., EPMA_TAPP_%s_Review%s_RD.xlsx)' % (
        VERSION, mode_tag, VERSION, mode_tag)

    # ---- what this is
    title('EPMA TAPP review workbook' + (' — %s' % mode if mode else ''))
    para('This workbook asks you to review the EPMA TAPP: the list of metadata items that should be '
         'reported when an EPMA procedure is registered, and when a measurement session run under that '
         'procedure is documented. It has four tabs: Instructions (this tab), Reviewer information, '
         'EPMA TAPP, and Registration Metadata. Please review the %d metadata items on the EPMA TAPP tab. '
         'The Registration Metadata tab lists %d items that are set by registration requirements; they '
         'need no review, though you may comment on them in "Other comments".'
         % (len(review_fields), len(reg_fields)))
    if mode:
        para('This copy covers %s only: the items that apply to that mode, %d of the %d in the full '
             'TAPP. "#" keeps the TAPP\'s own numbering, so gaps in the numbers are items that belong '
             'to other modes.' % (mode, len(fields), total_fields))
    para('Every item is described for two records. The procedure record is the standing recipe, '
         'registered once and cited by DOI. The session record describes one measurement session run '
         'under that procedure; a session may cover several samples. The procedure record is filled in '
         'once per analytical setup, typically from a representative session. A session record is filled '
         'in when data from that session are submitted, not for every session run on the instrument.')

    # ---- what to do
    head('What to do')
    pairs([('Step 1', 'Fill in the "Reviewer information" tab. Answers go in the "Your answer" column; '
                      'the questions are locked.', None, None),
           ('Step 2', 'On the "EPMA TAPP" tab, put all your feedback in the six shaded comment columns: '
                      '"Needed?", "Structure OK?" (twice), "Tier OK?" (twice) and "Other comments". '
                      'Each one judges the information column on its left; "Other comments" takes '
                      'everything else. Leave a cell blank if you agree or have no view. All other '
                      'columns are locked.', None, None),
           ('Step 3', 'Follow the "Review criteria" at the bottom of this tab. Provide review that clearly '
                      'presents your suggested change and reasons.', None, None),
           ('Step 4', 'When you have finished, save the workbook as %s and email it to info@astromat.org. '
                      'The deadline for this round of review is %s.' % (return_name, DEADLINE), None, None)],
          fill=STEP, label_size=22)

    # ---- the columns
    head('The EPMA TAPP tab: column explanations')
    info_names = ['#', 'Metadata item', 'Description', 'Example (NOT recommendations)',
                  'Structure — procedure record', 'Procedure-Level Tier', 'Structure — session record',
                  'Session-Level Tier'] + list(modes)
    para('The tab has %d information columns and 6 comment columns. The information columns are: %s. The '
         'comment columns are: "Needed?", "Structure OK?" (twice), "Tier OK?" (twice) and "Other '
         'comments". Each comment column sits immediately to the right of the information column it '
         'judges and has the same colour: purple for the metadata item, blue for the procedure record, '
         'yellow for the session record.' % (n_info, ', '.join('"%s"' % x for x in info_names)))
    for jj, h in enumerate(('Column', 'Meaning'), start=1):
        c = lg.cell(row=r, column=jj, value=h)
        c.font = Font(bold=True, size=18); c.fill = PAGE; c.border = BOX
        c.alignment = Alignment(horizontal='center', vertical='center')
    lg.row_dimensions[r].height = 30; r += 1
    guide = [
        ('#', 'The item\'s number in the TAPP: group number, then position in the group.', '44546A'),
        ('Metadata item', 'The name of the metadata item.', FAMILY['item']['dark']),
        ('Description', 'What the item records.', '44546A'),
        ('Example (NOT recommendations)', 'An illustration of the item filled in, from a fictional worked '
         'example (an invented EPMA session written up as a methods section). It shows the shape of an '
         'answer. It is not recommended content, and the laboratory, samples and results are invented. '
         'Samples are CC-1 and CC-2; analysis points are named by mineral and number (ol3 = the third '
         'olivine point, ap2 = the second apatite point); maps are Map 1 and Map 2.', '44546A'),
        ('Structure — procedure record', 'How the procedure states the item: one entry, one entry for each '
         'entry of a list, or a list of its own. The session record follows it, as shown in the session '
         'structure column.', FAMILY['procedure']['dark']),
        ('Procedure-Level Tier', 'Basic = mandatory for procedure registration. Advanced = recommended but '
         'optional. N/A = not part of the procedure: it can only be known once a session has run.',
         FAMILY['procedure']['dark']),
        ('Structure — session record', 'How one measurement session reports the item.',
         FAMILY['session']['dark']),
        ('Session-Level Tier', 'Read-Only = must follow the entry given in the procedure. Editable = may vary '
         'within the range the procedure allows; the session records what was used. Basic = the entry '
         'comes from the session itself and must be reported. Advanced = the entry comes from the session '
         'itself, and reporting it is recommended.', FAMILY['session']['dark']),
    ]
    if modes:
        guide.append((', '.join(modes), 'The analytical modes of EPMA. A tick means the item applies to '
                      'that mode.', '525252'))
    pairs([(a, b, None, c) for a, b, c in guide])

    head('The EPMA TAPP tab: "Tiers"')
    pairs([('Basic', 'Required', 'C00000', None), ('Advanced', 'Recommended', '375623', None),
           ('Read-Only', 'Fixed by the procedure. Session-level only.', '0070C0', None),
           ('Editable', 'Adjustable within the bounds the procedure sets. Session-level only.', '7030A0', None),
           ('N/A', 'Not applicable at that level. Procedure-level only.', '808080', None)])

    head('The EPMA TAPP tab: "Structure"')
    para('Each structure cell is an instruction, written in one of these patterns. An entry is one complete '
         'answer to the item, in whatever form its Description asks for: a number with its unit, a name, a '
         'choice, or a short statement with several parts (for example a reference material with its source '
         'and a citation). The parts of one entry are given together, not as separate entries. A list is '
         'made of entries too, one per member, and other items refer to its members as entries.')
    pairs([('Give one entry for each X.', 'One entry for every entry in the list X. "For the procedure" or "for '
            'the session" means nothing repeats: give a single entry. "One or more entries" means several may '
            'be given, but nothing else repeats over them.', None, None),
           ('For each A, give one entry for each B.', 'Two independent lists crossed: one entry for every '
            'pairing, with A as the outer list. Analytical Precision reads: for each "Secondary Reference '
            'Material", give one entry for each "Reported Variable".', None, None),
           ('For each A, give one entry for each of its B entries.', 'Nesting: B exists only inside A. For each '
            '"Sample Name", give one entry for each of its "Sampling Unit Name" entries — analysis points exist '
            'only inside a sample.', None, None),
           ('Give one entry; each "Sample Name" entry in a session inherits it.', 'Procedure column only. Samples, '
            'analysis points exist only in a session, so the procedure cannot list them: '
            'it gives one entry, and every sample (or point) in a session takes it.', None, None),
           ('List the …', 'The row is the list itself. Other rows are given for each of its entries, so it has '
            'to be filled in first.', None, None)])
    para('Names in quotes are lists, each defined by one row of the TAPP; the number after a name is that '
         'row\'s "#":')
    pairs([('Sample Name', 'the samples analysed in the session', None, None),
           ('Sampling Unit Name', 'the analysis points or map areas, per sample', None, None),
           ('Target Material', 'the material types the procedure analyses (row "Target Material")', None, None),
           ('Target Species', 'the elements the procedure determines (row "Target Species")', None, None),
           ('Monitored Element', 'what is measured to determine those elements (row "Monitored Elements")', None, None),
           ('Reported Variable', 'the reported quantities (row "Reported Variables and Units")', None, None),
           ('Secondary Reference Material', 'the secondary standards (row "Secondary Reference Materials")', None, None)])
    if REVIEW_NOTE:
        para('Cells shaded pale yellow in the procedure structure column are the ones we are least sure of. '
             'Hover over a shaded cell for what the literature shows.')

    head('Review criteria per column')
    pairs([('Needed?', 'Is this item relevant to EPMA as you practise it? Is it one thing, and distinct from '
            'other items? Is it correctly named?', FAMILY['item']['dark'], None),
           ('Structure OK? (procedure)', 'Does the procedure need one entry, one entry for each entry of a '
            'list (for example, per "Target Material"), or a list of its own?', FAMILY['procedure']['dark'], None),
           ('Tier OK? (procedure)', 'Can it be known before any session? If not: "N/A". If so, is it '
            'essential for judging data quality or reproducing the work? Then: "Basic"; otherwise: "Advanced".',
            FAMILY['procedure']['dark'], None),
           ('Structure OK? (session)', 'Does a session report it once, or once per sample, per analysis '
            'point, per element…?', FAMILY['session']['dark'], None),
           ('Tier OK? (session)', 'Must the session copy the procedure\'s entry unchanged (Should be fixed by '
            'the procedure, i.e. "Read-Only")? May it vary within the procedure\'s bounds (Should be adjustable, '
            'i.e. "Editable")? Or is it known only from the session ("Basic"/"Advanced")?',
            FAMILY['session']['dark'], None),
           ('Other comments', 'Anything else: the wording of Description or Example, missing items, wrong or '
            'missing modes, suggestions on the TAPP\'s structure.', FAMILY['comment']['dark'], None)])
    lg.column_dimensions['A'].width = WA
    lg.column_dimensions['B'].width = WB
    lg.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    lg.page_setup.orientation = 'landscape'
    lg.page_setup.fitToWidth = 1
    lg.page_setup.fitToHeight = 0

    # ============================== sheet 2: the table ==============================
    # ============================== sheet 2: reviewer information ======================
    # Added 2026-10-26. Free-text answers; the cells to fill are shaded like the review columns.
    # Completed workbooks carry personal data: collect them by email, never commit them to the
    # public repository.
    rv = wb.create_sheet('Reviewer information')
    rv.sheet_view.showGridLines = False
    t = rv.cell(row=1, column=1, value='About you')
    t.font = Font(bold=True, size=14)
    n = rv.cell(row=2, column=1, value='Please fill in the shaded cells. All answers are free text. '
                'Your details are used only to follow up on your review and to help us evaluate '
                'coverage of this review group; they are not published.')
    n.font = Font(italic=True, size=10, color='595959')
    n.alignment = Alignment(wrap_text=True, vertical='top')
    rv.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2)
    rv.row_dimensions[2].height = 30
    for j, h in enumerate(('Question', 'Your answer'), start=1):
        c = rv.cell(row=4, column=j, value=h)
        c.font = Font(bold=True, size=11, color='FFFFFF'); c.fill = HEAD_FILL; c.border = BORDER
    for i, q in enumerate(REVIEWER_QUESTIONS, start=5):
        a = rv.cell(row=i, column=1, value=q)
        a.font = Font(bold=True, size=14); a.alignment = Alignment(wrap_text=True, vertical='top')
        a.fill = LEGEND_FILL; a.border = BORDER
        b = rv.cell(row=i, column=2)
        b.fill = PatternFill('solid', fgColor=FAMILY['comment']['pale'])
        b.font = Font(size=14)
        b.alignment = Alignment(wrap_text=True, vertical='top'); b.border = BORDER
        b.protection = UNLOCKED
        rv.row_dimensions[i].height = est_height([(q, 44)], 14, minimum=30)
    rv.column_dimensions['A'].width = 44
    rv.column_dimensions['B'].width = 70
    protect(rv)

    def table(sheet_title, fields, placeholder=None):
        """One TAPP table: EPMA TAPP (items to review) or Registration Metadata."""
        ws = wb.create_sheet(sheet_title)
        ws.sheet_view.showGridLines = False
        spec = colspec(modes)
        ncols = len(spec)
        idx = {c[0]: i + 1 for i, c in enumerate(spec)}
        hdr = 1
        for j, (key, header, width, fam, kind, opts) in enumerate(spec, start=1):
            c = ws.cell(row=hdr, column=j, value=header)
            f = FAMILY.get(fam)
            if kind == 'review':
                c.fill = PatternFill('solid', fgColor=f['mid'])
                c.font = Font(bold=True, size=11, color='1F2A44')
            else:
                c.fill = PatternFill('solid', fgColor=f['dark'] if f else '44546A')
                c.font = Font(bold=True, size=11, color='FFFFFF')
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = BORDER
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.row_dimensions[hdr].height = 40

        row = hdr + 1
        flagged = []
        sections = sorted({f['sec'] for f in fields})
        for si, sec in enumerate(sections):
            members_all = [f for f in fields if f['sec'] == sec]
            if si:                                   # one blank row between groups
                row += 1
            for j in range(1, ncols + 1):
                cell = ws.cell(row=row, column=j); cell.fill = SECTION_FILL; cell.border = BORDER
            ws.cell(row=row, column=1, value=str(sec)).font = Font(bold=True, size=12)
            sname = members_all[0]['secname']
            ws.cell(row=row, column=2, value='Group %d: %s' % (sec, SECTION_SHORT.get(sname, sname))).font = Font(bold=True, size=12)
            ws.row_dimensions[row].height = 22
            row += 1
            if sec == 1 and placeholder:           # Group 1's registration rows live on their own tab
                first, last = placeholder[0]['num'], placeholder[-1]['num']
                ws.cell(row=row, column=1, value='%s–%s' % (first, last)).alignment = Alignment(vertical='top', wrap_text=True)
                txt = ('Items %s–%s are on the "Registration Metadata" tab. They are set by registration '
                       'requirements and need no review.' % (first, last))
                c = ws.cell(row=row, column=2, value=txt)
                c.font = Font(italic=True, size=11, color='595959')
                c.alignment = Alignment(wrap_text=True, vertical='top')
                for jj in range(1, ncols + 1):
                    ws.cell(row=row, column=jj).fill = LEGEND_FILL
                    ws.cell(row=row, column=jj).border = BORDER
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=idx['example'])
                ws.row_dimensions[row].height = 32
                row += 1

            # Rows are grouped by what they repeat over and by mode; the groups carry no
            # header of their own, only a rule above the first row of each.
            # GROUP_ROWS = False keeps the TAPP's own order (one "group" per field, no rules).
            groups = collections.OrderedDict()
            for f in members_all:
                if not GROUP_ROWS:
                    groups[(f['num'], f['modes'])] = [f]
                    continue
                sess, _ = phrase(f['name'], f['key'], 'session')
                groups.setdefault((sess.split(' per ', 1)[-1] if ' per ' in sess else sess,
                                   f['modes']), []).append(f)

            for gi, ((per, modepat), members) in enumerate(groups.items()):
                for f in members:
                    proc, flag = phrase(f['name'], f['key'], 'procedure')
                    if f['C'] == 'N/A':
                        proc, flag = '—', False
                    # Projection (conventions 7.3.3) is settled; flag only the cells with an open
                    # question attached, i.e. those that carry a review note.
                    flag = flag and f['name'] in REVIEW_NOTE
                    sess_f, _ = phrase(f['name'], f['key'], 'session')
                    if flag:
                        flagged.append(f['name'])
                    values = {'num': f['num'], 'name': f['name'], 'prose': PROSE[f['name']],
                              'example': example(f['name'], mode),
                              'pstruct': proc, 'ptier': f['C'],
                              'sstruct': sess_f, 'stier': f['D']}
                    for k, ch in zip(modes, f['modes']):
                        values['mode:' + k] = '✓' if ch == 'Y' else '–'
                    top = Side(style='thin', color='7F7F7F') if (GROUP_ROWS and f is members[0] and gi) else THIN
                    for j, (key, header, width, fam, kind, opts) in enumerate(spec, start=1):
                        cell = ws.cell(row=row, column=j, value=values.get(key))
                        cell.border = Border(left=THIN, right=THIN, top=top, bottom=THIN)
                        if kind == 'review':
                            cell.protection = UNLOCKED       # the only editable cells
                            cell.fill = PatternFill('solid', fgColor=FAMILY[fam]['pale'])
                            cell.alignment = Alignment(wrap_text=True, vertical='top')
                        elif key in ('ptier', 'stier'):
                            colour, bold = TIER_STYLE.get(cell.value, ('000000', False))
                            cell.font = Font(bold=bold, color=colour, size=11)
                            cell.alignment = Alignment(horizontal='center', vertical='top')
                        elif key.startswith('mode:'):
                            tick = cell.value == '✓'
                            cell.font = Font(bold=tick, size=11,
                                             color='375623' if tick else 'BFBFBF')
                            if tick:
                                cell.fill = TICK_FILL
                            cell.alignment = Alignment(horizontal='center', vertical='top')
                        elif key == 'name':
                            cell.font = Font(bold=True, size=11)
                            cell.alignment = Alignment(vertical='top', wrap_text=True)
                        else:
                            cell.alignment = Alignment(wrap_text=(key in ('prose', 'example', 'pstruct', 'sstruct')), vertical='top')
                    if f['name'] in REGISTRATION:      # judgement cells closed; Other comments open
                        for j, (key, header, width, fam, kind, opts) in enumerate(spec, start=1):
                            if kind == 'review' and opts:
                                c = ws.cell(row=row, column=j)
                                c.protection = Protection(locked=True)
                                c.fill = GREY_FILL
                    ws.row_dimensions[row].height = est_height(
                        [(values.get(k), w) for (k, h_, w, *_rest) in spec])
                    if flag:
                        c = ws.cell(row=row, column=idx['pstruct'])
                        c.fill = FLAG_FILL
                        note = REVIEW_NOTE.get(f['name'], '')
                        c.comment = Comment((note + NOTE_TAIL).strip(), 'TAPP review',
                                            height=130, width=330)
                    row += 1

        # Free text throughout: no drop-downs (RD, 2026-10-26).
        ws.freeze_panes = ws.cell(row=hdr + 1, column=3)
        ws.auto_filter.ref = 'A%d:%s%d' % (hdr, get_column_letter(ncols), row - 1)
        protect(ws)
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = 8
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '%d:%d' % (hdr, hdr)
        return flagged

    flagged = table('EPMA TAPP', review_fields, placeholder=reg_fields)
    flagged += table('Registration Metadata', reg_fields)

    # ============================== sheet 3: traceability ==========================
    tr = wb.create_sheet('Structure key (traceability)')
    tr.sheet_state = 'hidden'          # for the TAPP team, not reviewers (2026-10-26)
    tr.sheet_view.showGridLines = False
    for j, h in enumerate(['#', 'Metadata item', 'TAPP "Keyed By" (Column I)',
                           'Structure — procedure record', 'Structure — session record'], start=1):
        c = tr.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, size=11, color='FFFFFF'); c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    for i, f in enumerate(fields, start=2):
        proc, flag = phrase(f['name'], f['key'], 'procedure')
        if f['C'] == 'N/A':
            proc, flag = '—', False
        flag = flag and f['name'] in REVIEW_NOTE      # same rule as the main sheet
        sess, _ = phrase(f['name'], f['key'], 'session')
        for j, v in enumerate([f['num'], f['name'], f['key'], proc, sess], start=1):
            cell = tr.cell(row=i, column=j, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        if flag:
            tr.cell(row=i, column=4).fill = FLAG_FILL
    for col, w in zip('ABCDE', [7, 34, 40, 30, 32]):
        tr.column_dimensions[col].width = w
    tr.freeze_panes = 'A2'
    tr.auto_filter.ref = 'A1:E%d' % (len(fields) + 1)

    out = out_path(mode)
    wb.save(out)
    return out, len(fields), sorted(set(flagged))


if __name__ == '__main__':
    print('read   %s' % SRC)
    os.makedirs(OUT_DIR, exist_ok=True)
    for m in (None,) + MODES:
        out, n, flagged = build(m)
        print('wrote %s — %d fields, %d flagged' % (os.path.relpath(out, ROOT), n, len(flagged)))
    for fn in prune_other_versions():
        print('removed superseded %s' % fn)
